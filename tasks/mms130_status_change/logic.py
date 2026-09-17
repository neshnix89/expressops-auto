"""
Pure logic for mms130_status_change: read the input sheet, normalise
columns/values, and write the results workbook. No browser code here.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

# Accepted header spellings (lower-cased, spaces collapsed) → canonical key.
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "item": ("item number", "item", "item no", "pn", "part number", "part no",
             "itno", "article", "article number"),
    "lot": ("lot number", "lot", "lot no", "bano", "batch"),
    "warehouse": ("warehouse", "whs", "whlo", "wh"),
    "location": ("location", "loc", "whsl", "stock location"),
    "container": ("container", "camu"),
}
REQUIRED = ("item", "lot", "warehouse", "location")


@dataclass
class InputRow:
    row_no: int                     # sheet row number (header = 1)
    item: str
    lot: str
    warehouse: str
    location: str
    container: str = ""
    result: str = ""                # OK / DRY-RUN OK / ERROR / SKIPPED
    message: str = ""
    panel_b: dict[str, Any] = field(default_factory=dict)

    @property
    def label(self) -> str:
        return f"row {self.row_no}: {self.item} / lot {self.lot} @ {self.warehouse}/{self.location}"


def _norm_header(h: Any) -> str:
    return " ".join(str(h or "").replace("_", " ").split()).lower()


def _cell(v: Any) -> str:
    """Excel numbers come back as floats (12345.0) — keep them as clean text."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def map_headers(headers: list[Any]) -> dict[str, int]:
    """Return {canonical_key: column_index}. Raises ValueError if a required column is missing."""
    mapping: dict[str, int] = {}
    for idx, h in enumerate(headers):
        n = _norm_header(h)
        for key, aliases in COLUMN_ALIASES.items():
            if n in aliases and key not in mapping:
                mapping[key] = idx
    missing = [k for k in REQUIRED if k not in mapping]
    if missing:
        raise ValueError(
            f"Input is missing column(s): {', '.join(missing)}. "
            f"Found headers: {[str(h) for h in headers]}. "
            "Expected e.g. 'Item number', 'Lot number', 'Warehouse', 'Location'."
        )
    return mapping


def _raw_rows(path: Path, sheet: str | None) -> list[list[Any]]:
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            return [list(r) for r in csv.reader(f)]
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def read_input(path: Path, sheet: str | None = None) -> list[InputRow]:
    """Read the input sheet. Blank lines are skipped; incomplete lines are returned as SKIPPED."""
    raw = _raw_rows(path, sheet)
    if not raw:
        raise ValueError(f"{path.name} is empty")
    cols = map_headers(raw[0])
    rows: list[InputRow] = []
    for i, r in enumerate(raw[1:], start=2):
        get = lambda k: _cell(r[cols[k]]) if k in cols and cols[k] < len(r) else ""  # noqa: E731
        vals = {k: get(k) for k in COLUMN_ALIASES}
        if not any(vals.values()):
            continue
        row = InputRow(row_no=i, **vals)
        missing = [k for k in REQUIRED if not vals[k]]
        if missing:
            row.result, row.message = "SKIPPED", f"missing {', '.join(missing)}"
        rows.append(row)
    return rows


def write_results(rows: list[InputRow], out_path: Path, mode: str) -> Path:
    """Write a results workbook (falls back to CSV if openpyxl is unavailable)."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    header = ["Row", "Item number", "Lot number", "Warehouse", "Location",
              "Container", "Result", "Message", "Panel B snapshot"]
    data = [[r.row_no, r.item, r.lot, r.warehouse, r.location, r.container,
             r.result, r.message,
             "; ".join(f"{k}={v}" for k, v in r.panel_b.items())] for r in rows]
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        out_path = out_path.with_suffix(".csv")
        with out_path.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(data)
        return out_path

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    fills = {"OK": "C6EFCE", "DRY-RUN OK": "DDEBF7", "ERROR": "FFC7CE", "SKIPPED": "FFEB9C"}
    for d in data:
        ws.append(d)
        colour = fills.get(d[6])
        if colour:
            ws.cell(ws.max_row, 7).fill = PatternFill("solid", fgColor=colour)
    for col, width in zip("ABCDEFGHI", (6, 18, 18, 11, 14, 14, 12, 60, 80)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    meta = wb.create_sheet("Run")
    meta.append(["Mode", mode])
    meta.append(["Run at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    meta.append(["Rows", len(rows)])
    for k in ("OK", "DRY-RUN OK", "ERROR", "SKIPPED"):
        meta.append([k, sum(1 for r in rows if r.result == k)])
    wb.save(out_path)
    return out_path
